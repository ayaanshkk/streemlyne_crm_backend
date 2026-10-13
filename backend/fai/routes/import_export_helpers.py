"""
Import / Export helpers for StreemLyne CRM
──────────────────────────────────────────────────────────────────────────
Shared by company_routes.py and contact_routes.py. Small and dependency-
light on purpose: openpyxl is the only extra package needed (add it to
requirements.txt if it isn't already there — pricelist_routes.py may
already depend on it, in which case this is a no-op).

Two responsibilities:
    1. parse_uploaded_rows(file_storage)  -> list[dict]   (import)
    2. build_csv_response(fieldnames, rows, filename)     (export)

Import always works off the FIRST ROW as a header, matched
case-insensitively and with surrounding whitespace stripped, so
"Company Name", "company_name", and " company_name " all resolve to the
same key. Unrecognized columns are ignored rather than erroring, so a
slightly messy export-then-reimport round-trip doesn't break.

UPDATED: file type is now detected from the file's actual content (a
real .xlsx is a zip archive, so it starts with the bytes "PK") instead
of trusting the uploaded filename's extension. This fixes a real-world
case where Export always produces CSV content but the browser download
was named "companies.xlsx" — re-importing that exact file used to fail
with "File is not a zip file" because the extension said .xlsx while
the content was plain CSV text. Sniffing the content means a file's
real format is what decides how it's parsed, regardless of what it's
named.
"""

import csv
import io

import openpyxl
from werkzeug.datastructures import FileStorage

class ImportFileError(ValueError):
    """Raised for any problem reading/parsing the uploaded file."""
    pass


# Every .xlsx/.xlsm file is a zip archive under the hood, and every zip
# archive starts with these two bytes ("PK", after zip's original author
# Phil Katz). Checking this instead of the filename means a mislabeled
# file (e.g. CSV content saved with a .xlsx extension) still gets parsed
# correctly, and a genuine .xlsx still works no matter what it's named.
ZIP_MAGIC_BYTES = b"PK"


def _normalize_header(value):
    return (str(value) if value is not None else "").strip().lower().replace(" ", "_")


def _looks_like_xlsx(raw_bytes: bytes) -> bool:
    return raw_bytes[:2] == ZIP_MAGIC_BYTES


def _parse_xlsx_bytes(raw_bytes: bytes):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception as e:
        raise ImportFileError(f"Could not read .xlsx file: {e}")

    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = [_normalize_header(h) for h in next(rows_iter)]
    except StopIteration:
        return []

    results = []
    for raw_row in rows_iter:
        if raw_row is None or all(cell in (None, "") for cell in raw_row):
            continue
        row = {}
        for key, value in zip(header, raw_row):
            if not key:
                continue
            row[key] = "" if value is None else str(value).strip()
        results.append(row)
    return results


def _parse_csv_bytes(raw_bytes: bytes):
    try:
        text_stream = io.StringIO(raw_bytes.decode("utf-8-sig"))
    except Exception as e:
        raise ImportFileError(f"Could not read .csv file: {e}")

    reader = csv.DictReader(text_stream)
    results = []
    for raw_row in reader:
        if not any((v or "").strip() for v in raw_row.values()):
            continue
        row = {
            _normalize_header(k): (v or "").strip()
            for k, v in raw_row.items() if k
        }
        results.append(row)
    return results


def parse_uploaded_rows(file_storage: FileStorage):
    """
    Reads a .csv or .xlsx upload and returns a list of dicts keyed by
    normalized header name (lowercase, spaces -> underscores). Blank rows
    (no non-empty cell) are skipped.

    File type is determined by sniffing the actual bytes (zip signature
    = xlsx, otherwise treated as CSV/text) rather than trusting the
    uploaded filename's extension — see module docstring for why.
    """
    filename = (file_storage.filename or "").lower()
    raw_bytes = file_storage.read()

    if not raw_bytes:
        raise ImportFileError("The uploaded file is empty")

    if _looks_like_xlsx(raw_bytes):
        return _parse_xlsx_bytes(raw_bytes)

    if filename.endswith((".csv", ".xlsx", ".xls")) or filename == "":
        # Not a zip signature, but named like a spreadsheet (or has no
        # extension at all) — most likely CSV content, possibly
        # mislabeled. Try parsing it as CSV rather than rejecting it
        # outright based on the extension alone.
        return _parse_csv_bytes(raw_bytes)

    raise ImportFileError("Unsupported file type — please upload a .csv or .xlsx file")

def apply_header_aliases(rows, alias_map):
    """
    Remaps parsed rows so alternate/synonym column headers resolve to
    the canonical field name a route actually expects.

    `alias_map` is {canonical_field: [alias1, alias2, ...]}, where each
    alias is written pre-normalized (lowercase, spaces -> underscores) —
    e.g. "contact_number" not "Contact Number".

    For each row and each canonical field:
      - if the canonical header itself is present and non-empty, it wins
        as-is (no override).
      - otherwise, aliases are tried in order; the first one with a
        non-empty value in that row is copied into the canonical key.
      - columns that match no alias for any canonical field are left in
        the row untouched (so callers can still see extras) but are
        simply never read by the route — this is the "ignore" behavior.

    This only fills in canonical keys; it never deletes existing columns.
    """
    normalized_alias_map = {
        canonical: [a.strip().lower().replace(" ", "_") for a in aliases]
        for canonical, aliases in alias_map.items()
    }

    mapped_rows = []
    for row in rows:
        new_row = dict(row)
        for canonical, aliases in normalized_alias_map.items():
            if (row.get(canonical) or "").strip():
                continue  # canonical header already present and filled
            for alias in aliases:
                value = (row.get(alias) or "").strip()
                if value:
                    new_row[canonical] = value
                    break
        mapped_rows.append(new_row)
    return mapped_rows


def build_csv_response(fieldnames, rows, filename):
    """
    Builds a Flask Response streaming a CSV attachment. `rows` is a list
    of dicts; only keys in `fieldnames` are written, in that order.
    """
    from flask import Response

    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=fieldnames, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow(row)

    return Response(
        buffer.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

def build_xlsx_response(fieldnames, rows, filename):
    """
    Builds a Flask Response streaming a real .xlsx attachment, built
    with openpyxl (same library this module already uses for reading).
    """
    from flask import Response

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.append(fieldnames)
    for row in rows:
        sheet.append([row.get(f, "") for f in fieldnames])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
