"""
backend/routes/import_export_helpers.py

Small CSV / XLSX helpers shared by company_routes.py and contact_routes.py.

    parse_uploaded_rows(file_storage) -> list[dict]
        Reads an uploaded .csv or .xlsx (first sheet). The first row is the
        header; header names are normalised to lower_snake_case
        ("Company Name" -> "company_name"). Blank rows are skipped. Every
        value comes back as a stripped string ('' for empty cells).

    apply_header_aliases(rows, aliases) -> list[dict]
        Renames alternative column headers to canonical ones, e.g. with
        {"phone": ["mobile", "phone_no"]} a column called "Mobile" becomes
        "phone". If a row has several aliases for the same field, the first
        non-empty value wins. Unknown columns pass through unchanged.

    build_csv_response(fields, rows, filename)  -> flask.Response
    build_xlsx_response(fields, rows, filename) -> flask.Response
        `fields` is the ordered list of keys to export, `rows` is a list of
        dicts.

    ImportFileError
        Raised for anything wrong with the uploaded file; the routes turn it
        into a 400.
"""

import csv
import io
import re

from flask import Response

MAX_IMPORT_ROWS = 5000


class ImportFileError(Exception):
    """The uploaded file can't be read or has an unusable shape."""


def _normalise_header(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"[^\w]+", "_", text)
    return text.strip("_")


def _cell_to_str(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # 411001.0 -> "411001" (Excel stores numbers as floats)
    return str(value).strip()


def _rows_to_dicts(header, body_rows):
    keys = [_normalise_header(h) for h in header]
    if not any(keys):
        raise ImportFileError("The first row must contain column headers")

    out = []
    for raw in body_rows:
        values = [_cell_to_str(v) for v in raw]
        if not any(values):
            continue  # skip blank rows
        values += [""] * (len(keys) - len(values))
        out.append({k: v for k, v in zip(keys, values) if k})
        if len(out) > MAX_IMPORT_ROWS:
            raise ImportFileError(f"Too many rows (limit is {MAX_IMPORT_ROWS})")
    return out


def parse_uploaded_rows(file_storage):
    filename = (getattr(file_storage, "filename", "") or "").lower()
    if not filename:
        raise ImportFileError("No file selected")

    data = file_storage.read()
    if not data:
        raise ImportFileError("The uploaded file is empty")

    if filename.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")  # Excel "CSV" exports are often cp1252
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if not all_rows:
            raise ImportFileError("The uploaded file is empty")
        return _rows_to_dicts(all_rows[0], all_rows[1:])

    if filename.endswith((".xlsx", ".xlsm")):
        try:
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            all_rows = list(ws.iter_rows(values_only=True))
        except ImportFileError:
            raise
        except Exception as e:
            raise ImportFileError(f"Could not read the Excel file: {e}")
        if not all_rows:
            raise ImportFileError("The uploaded file is empty")
        return _rows_to_dicts(all_rows[0], all_rows[1:])

    raise ImportFileError("Unsupported file type. Upload a .csv or .xlsx file")


def apply_header_aliases(rows, aliases):
    lookup = {}
    for canonical, names in aliases.items():
        lookup[_normalise_header(canonical)] = canonical
        for name in names:
            lookup[_normalise_header(name)] = canonical

    out = []
    for row in rows:
        new_row = {}
        for key, value in row.items():
            canonical = lookup.get(key, key)
            if canonical not in new_row or (not new_row[canonical] and value):
                new_row[canonical] = value
        out.append(new_row)
    return out


def build_csv_response(fields, rows, filename):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(fields)
    for row in rows:
        writer.writerow(["" if row.get(f) is None else row.get(f) for f in fields])
    # utf-8-sig so Excel opens non-ASCII text correctly
    body = buf.getvalue().encode("utf-8-sig")
    return Response(
        body,
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def build_xlsx_response(fields, rows, filename):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(list(fields))
    for row in rows:
        ws.append(["" if row.get(f) is None else row.get(f) for f in fields])
    out = io.BytesIO()
    wb.save(out)
    return Response(
        out.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )